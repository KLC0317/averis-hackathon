from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl
import pymupdf as fitz
import pytest
from docx import Document

from cleardraft.core import DocumentRole, ValueStatus, extract_fields
from cleardraft.pipeline import verify_case
from cleardraft.readers import read_document, read_docx, read_pdf, read_txt, read_xlsx


BUNDLE_ATTACHMENTS = Path("sdoc-hackathon-bundle/attachments")
BUNDLE_INBOX = Path("sdoc-hackathon-bundle/inbox")


def _make_sample_text(role: str = "SI") -> str:
    heading = "SHIPPING INSTRUCTION" if role == "SI" else "BILL OF LADING (DRAFT)"
    return f"""{heading}
Shipper: Ocean Cargo Logistics Ltd
Consignee: Global Import Corp
Notify Party: Global Import Corp
Port of Loading: Port Klang
Port of Discharge: Jebel Ali
Container Count: 2 x 40HC
Gross Weight: 25400 KG
"""


def _make_synthetic_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_scanned_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, (0, 0, 100, 100), 0)
    page.insert_image(page.rect, pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data


def _make_synthetic_docx(text: str) -> bytes:
    doc = Document()
    lines = text.strip().splitlines()
    doc.add_paragraph(lines[0])
    table = doc.add_table(rows=len(lines) - 1, cols=2)
    for idx, line in enumerate(lines[1:]):
        if ":" in line:
            label, val = line.split(":", 1)
            table.rows[idx].cells[0].text = label.strip()
            table.rows[idx].cells[1].text = val.strip()
        else:
            table.rows[idx].cells[0].text = line.strip()
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def _make_synthetic_xlsx(text: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for line in text.strip().splitlines():
        if ":" in line:
            label, val = line.split(":", 1)
            ws.append([label.strip(), val.strip()])
        else:
            ws.append([line.strip()])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# 1. TXT path
def test_txt_path_exercised() -> None:
    si_text = _make_sample_text("SI")
    bl_text = _make_sample_text("BL")

    si_res = read_document(si_text.encode("utf-8"), "case_SI.txt")
    assert si_res.detected_format == "txt"
    assert len(si_res.blocks) >= 7
    assert si_res.blocks[0].locator["start_line"] == 1
    assert si_res.error is None

    si_fields = extract_fields(si_res.text, DocumentRole.SI)
    assert si_fields["shipper"].value_status == ValueStatus.PRESENT
    assert si_fields["shipper"].canonical_value == "ocean cargo logistics ltd"

    email = {"email_id": "txt_test", "subject": "Compare docs", "attachments": ["attachments/si.txt", "attachments/bl.txt"]}
    docs = [
        {"filename": "si.txt", "bytes": si_text.encode("utf-8")},
        {"filename": "bl.txt", "bytes": bl_text.encode("utf-8")},
    ]
    res = verify_case(docs, email, "BL_COMPARISON")
    sub = res.to_submission()
    assert res.verification == "MATCH"
    assert sub["status"] == "OK"
    assert sub["has_defect"] is False


# 2. Native PDF path
def test_native_pdf_path_exercised() -> None:
    pdf_bytes = _make_synthetic_pdf(_make_sample_text("BL"))
    pdf_res = read_document(pdf_bytes, "draft_BL.pdf")
    assert pdf_res.detected_format == "pdf"
    assert len(pdf_res.blocks) == 1
    assert pdf_res.blocks[0].locator["kind"] == "page"
    assert pdf_res.blocks[0].locator["page"] == 1
    assert pdf_res.ocr_used is False
    assert pdf_res.error is None

    bl_fields = extract_fields(pdf_res.text, DocumentRole.DRAFT_BL)
    assert bl_fields["shipper"].canonical_value == "ocean cargo logistics ltd"
    assert bl_fields["container_count"].canonical_value == "2"

    if BUNDLE_ATTACHMENTS.exists():
        real_pdf = (BUNDLE_ATTACHMENTS / "email_059_BL.pdf").read_bytes()
        real_res = read_document(real_pdf, "email_059_BL.pdf")
        assert real_res.detected_format == "pdf"
        assert len(real_res.blocks) >= 1
        assert real_res.error is None


# 3. Scanned PDF / OCR fallback path
def test_scanned_pdf_and_ocr_fallback_path_exercised() -> None:
    scanned_bytes = _make_scanned_pdf()
    res = read_document(scanned_bytes, "scanned_si.pdf")
    assert res.detected_format == "pdf"
    assert "page_1_has_no_text" in res.warnings
    assert "ocr_required" in res.warnings
    assert res.ocr_used is False
    assert len(res.blocks) == 0

    si_txt = _make_sample_text("SI").encode("utf-8")
    email = {"email_id": "scan_test", "subject": "Compare docs", "attachments": ["attachments/si.txt", "attachments/bl.pdf"]}
    docs = [
        {"filename": "si.txt", "bytes": si_txt},
        {"filename": "bl.pdf", "bytes": scanned_bytes},
    ]
    verification = verify_case(docs, email, "BL_COMPARISON")
    sub = verification.to_submission()
    assert verification.verification == "NEEDS_REVIEW"
    assert sub["status"] == "NEEDS_REVIEW"
    assert sub["review_reason"] == "unreadable"
    assert any("ocr_required" in issue for issue in verification.issues)


# 4. DOCX path
def test_docx_path_exercised() -> None:
    docx_bytes = _make_synthetic_docx(_make_sample_text("BL"))
    docx_res = read_document(docx_bytes, "draft.docx")
    assert docx_res.detected_format == "docx"
    assert len(docx_res.blocks) >= 7
    assert docx_res.error is None
    locators = [b.locator["kind"] for b in docx_res.blocks]
    assert "paragraph" in locators or "table_cell_row" in locators

    bl_fields = extract_fields(docx_res.text, DocumentRole.DRAFT_BL)
    assert bl_fields["shipper"].canonical_value == "ocean cargo logistics ltd"
    assert bl_fields["gross_weight_kg"].canonical_value == "25400"

    if BUNDLE_ATTACHMENTS.exists():
        real_docx = (BUNDLE_ATTACHMENTS / "email_055_BL.docx").read_bytes()
        real_res = read_document(real_docx, "email_055_BL.docx")
        assert real_res.detected_format == "docx"
        assert len(real_res.blocks) >= 5
        assert real_res.error is None


# 5. XLSX path
def test_xlsx_path_exercised() -> None:
    xlsx_bytes = _make_synthetic_xlsx(_make_sample_text("SI"))
    xlsx_res = read_document(xlsx_bytes, "si.xlsx")
    assert xlsx_res.detected_format == "xlsx"
    assert len(xlsx_res.blocks) >= 7
    assert xlsx_res.error is None
    assert xlsx_res.blocks[0].locator["kind"] == "worksheet_cells"

    si_fields = extract_fields(xlsx_res.text, DocumentRole.SI)
    assert si_fields["shipper"].canonical_value == "ocean cargo logistics ltd"
    assert si_fields["port_of_discharge"].canonical_value == "jebel ali"

    if BUNDLE_ATTACHMENTS.exists():
        real_xlsx = (BUNDLE_ATTACHMENTS / "email_005_SI.xlsx").read_bytes()
        real_res = read_document(real_xlsx, "email_005_SI.xlsx")
        assert real_res.detected_format == "xlsx"
        assert len(real_res.blocks) >= 5
        assert real_res.error is None


# 6. Corrupt documents path
def test_corrupt_files_path_exercised() -> None:
    corrupt_pdf = b"%PDF-broken-header-truncated-bytes"
    res_pdf = read_document(corrupt_pdf, "broken.pdf")
    assert res_pdf.detected_format == "pdf"
    assert res_pdf.error is not None
    assert "pdf_read_failed" in res_pdf.error

    corrupt_docx = b"PK\x03\x04not-a-valid-docx-archive"
    res_docx = read_document(corrupt_docx, "broken.docx")
    assert res_docx.detected_format == "docx"
    assert res_docx.error is not None
    assert "docx_read_failed" in res_docx.error

    corrupt_xlsx = b"PK\x03\x04not-a-valid-xlsx-archive"
    res_xlsx = read_document(corrupt_xlsx, "broken.xlsx")
    assert res_xlsx.detected_format == "xlsx"
    assert res_xlsx.error is not None
    assert "xlsx_read_failed" in res_xlsx.error

    # Corrupt document in pipeline leads to FAILED processing and unreadable reason
    email = {"email_id": "corrupt_test", "subject": "Compare docs", "attachments": ["attachments/si.txt", "attachments/broken_BL.pdf"]}
    docs = [
        {"filename": "si.txt", "bytes": _make_sample_text("SI").encode("utf-8")},
        {"filename": "broken_BL.pdf", "bytes": corrupt_pdf},
    ]
    res = verify_case(docs, email, "BL_COMPARISON")
    sub = res.to_submission()
    assert res.processing == "FAILED"
    assert res.verification == "NEEDS_REVIEW"
    assert sub["status"] == "NEEDS_REVIEW"
    assert sub["review_reason"] == "unreadable"


# 7. Missing source paths
def test_missing_source_paths_exercised() -> None:
    si_bytes = _make_sample_text("SI").encode("utf-8")
    bl_bytes = _make_sample_text("BL").encode("utf-8")

    # Path 7a: Zero attachments provided for comparison request
    res_zero = verify_case([], {"email_id": "e_zero", "attachments": []}, "BL_COMPARISON")
    sub_zero = res_zero.to_submission()
    assert sub_zero["status"] == "NEEDS_REVIEW"
    assert sub_zero["review_reason"] == "missing_attachment"

    # Path 7b: Only SI provided (draft BL missing)
    res_no_bl = verify_case([{"filename": "si.txt", "bytes": si_bytes}], {"email_id": "e_no_bl", "attachments": ["si.txt"]}, "BL_COMPARISON")
    sub_no_bl = res_no_bl.to_submission()
    assert sub_no_bl["status"] == "NEEDS_REVIEW"
    assert sub_no_bl["review_reason"] == "missing_attachment"

    # Path 7c: Only BL provided (SI missing)
    res_no_si = verify_case([{"filename": "bl.txt", "bytes": bl_bytes}], {"email_id": "e_no_si", "attachments": ["bl.txt"]}, "BL_COMPARISON")
    sub_no_si = res_no_si.to_submission()
    assert sub_no_si["status"] == "NEEDS_REVIEW"
    assert sub_no_si["review_reason"] == "missing_attachment"

    # Path 7d: Wrong document type (Commercial Invoice provided instead of draft BL)
    inv_bytes = b"COMMERCIAL INVOICE\nInvoice No: 12345\nTotal: USD 50000\nShipper: Ocean Cargo\n"
    res_wrong = verify_case([
        {"filename": "si.txt", "bytes": si_bytes},
        {"filename": "bl.txt", "bytes": inv_bytes},
    ], {"email_id": "e_wrong", "attachments": ["si.txt", "bl.txt"]}, "BL_COMPARISON")
    sub_wrong = res_wrong.to_submission()
    assert sub_wrong["status"] == "NEEDS_REVIEW"
    assert sub_wrong["review_reason"] == "wrong_doc_type"

    # Path 7e: Missing field value in valid document
    si_missing_fld = """SHIPPING INSTRUCTION
Shipper: Ocean Cargo Logistics Ltd
Consignee: Global Import Corp
Notify Party: Global Import Corp
Port of Loading: Port Klang
Port of Discharge:
Container Count: 2 x 40HC
Gross Weight: 25400 KG
""".encode("utf-8")
    res_miss_val = verify_case([
        {"filename": "si.txt", "bytes": si_missing_fld},
        {"filename": "bl.txt", "bytes": bl_bytes},
    ], {"email_id": "e_miss_val", "attachments": ["si.txt", "bl.txt"]}, "BL_COMPARISON")
    sub_miss_val = res_miss_val.to_submission()
    assert sub_miss_val["status"] == "NEEDS_REVIEW"
    assert sub_miss_val["review_reason"] == "missing_value"
    assert "port_of_discharge" in res_miss_val.unresolved_fields
